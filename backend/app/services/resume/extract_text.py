from __future__ import annotations

import io
import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document


def extract_text_from_file(file_path: str | Path, content_type: str | None = None) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf" or (content_type and "pdf" in content_type):
        return _extract_pdf(path)
    if suffix == ".docx" or (content_type and "wordprocessingml" in (content_type or "")):
        return _extract_docx(path)
    if suffix == ".txt" or (content_type and content_type.startswith("text/")):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume type: {suffix or content_type}")


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if suffix == ".txt":
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"Unsupported resume type: {suffix}")


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
PHONE_RE = re.compile(
    r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{2,4}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,6}"
)
SECTION_HEADERS = {
    "summary": re.compile(r"^\s*(summary|profile|objective|about)\s*:?\s*$", re.I),
    "experience": re.compile(
        r"^\s*(experience|work experience|employment|professional experience)\s*:?\s*$", re.I
    ),
    "projects": re.compile(r"^\s*(projects|personal projects|key projects)\s*:?\s*$", re.I),
    "education": re.compile(r"^\s*(education|academic)\s*:?\s*$", re.I),
    "skills": re.compile(
        r"^\s*(skills|technical skills|core competencies|technologies)\s*:?\s*$", re.I
    ),
    "certifications": re.compile(r"^\s*(certifications?|licenses?)\s*:?\s*$", re.I),
    "achievements": re.compile(r"^\s*(achievements?|awards?|honors?)\s*:?\s*$", re.I),
}


def split_sections(text: str) -> dict[str, str]:
    lines = text.splitlines()
    sections: dict[str, list[str]] = {"_preamble": []}
    current = "_preamble"
    for line in lines:
        matched = None
        for name, pattern in SECTION_HEADERS.items():
            if pattern.match(line.strip()):
                matched = name
                break
        if matched:
            current = matched
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if v}
